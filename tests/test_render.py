"""End-to-end: example snapshot -> As-Built + Summary as PDF and Word."""

from pathlib import Path

import pytest

from sextant.builders.asbuilt import build_asbuilt
from sextant.builders.summary import build_summary
from sextant.doc import DeviceHeader, KpiRow, PageBreak
from sextant.pipeline import load_snapshot

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "snapshot.json"
pytestmark = pytest.mark.skipif(not EXAMPLE.exists(), reason="examples/snapshot.json not built")


@pytest.fixture(scope="module")
def snap():
    return load_snapshot(EXAMPLE)


def test_summary_one_page_per_device(snap):
    doc = build_summary(snap)
    assert sum(isinstance(b, DeviceHeader) for b in doc.blocks) == len(snap.devices)
    topo_pages = sum(1 for b in doc.blocks if getattr(b, "anchor", None) == "sum-topology")
    assert sum(isinstance(b, PageBreak) for b in doc.blocks) == len(snap.devices) + topo_pages
    for b in doc.blocks:
        if isinstance(b, KpiRow):
            assert 1 <= len(b.items) <= 6


def test_render_all_formats(snap, tmp_path):
    from docx import Document

    from sextant.render import render_all

    paths = render_all(snap, tmp_path, include_raw=True)
    assert len(paths) == 6 and all(p.stat().st_size > 10_000 for p in paths)
    asbuilt_docx = next(p for p in paths if p.name.endswith("As-Built.docx"))
    text = "\n".join(p.text for p in Document(asbuilt_docx).paragraphs)
    assert "Executive overview" in text
    for d in snap.devices:
        assert d.hostname in text


def test_no_secrets_leak(snap):
    from sextant.render.html import render_html

    html = render_html(build_asbuilt(snap, include_raw=True))
    for needle in ("$9$", "$1$", "ENC SH2", "BEGIN RSA PRIVATE KEY-----\nMII"):
        assert needle not in html


def _device_pages(snap, doc):
    """PDF pages each device's Summary page takes (from where its heading anchor lands)."""
    from weasyprint import HTML

    from sextant.doc import Heading
    from sextant.render.html import render_html

    pages = HTML(string=render_html(doc)).render().pages
    where = {a: i for i, p in enumerate(pages) for a in p.anchors}
    starts = [where[b.anchor] for b in doc.blocks if isinstance(b, Heading) and b.anchor.startswith("sum-")
              and b.anchor != "sum-topology" and b.anchor in where]
    assert len(starts) == len(snap.devices)
    return [b - a for a, b in zip(starts, [*starts[1:], len(pages)], strict=True)]


def test_summary_device_pages_fit_on_one_pdf_page(snap):
    assert set(_device_pages(snap, build_summary(snap))) == {1}


def test_summary_long_vlan_list_is_cut_to_fit():
    """A switch with many VLANs still gets one page: the VLAN table is cut to the room left after the
    findings and glance facts, down to a pointer to the As-Built when nothing fits."""
    from sextant.doc import Paragraph
    from sextant.model import Device, Finding, Interface, PortMode, Severity, Snapshot, Vlan

    def switch(n_findings):
        return Device(
            hostname="big-sw", role="switch", vendor="Cisco", model="C9300-48P", os_name="IOS-XE", os_version="17.9",
            mgmt_ip="10.0.0.1", uptime="214 day(s), 6 hour(s), 31 minute(s), 9 second(s)",
            vlans=[Vlan(id=i, name=f"VLAN-NAME-{i}", svi_ipv4=[f"10.{i}.0.1/24"]) for i in range(1, 60)],
            interfaces=[Interface(name=f"Port-channel{i}", kind="lag", mode=PortMode.trunk,
                                  lag_members=["Te1/1/1", "Te1/1/2"]) for i in range(1, 9)],
            findings=[Finding(check_id=f"x{i}", severity=Severity.medium, title="A finding title " * 6,
                              recommendation="Do this and then that " * 6) for i in range(n_findings)])

    vlan_tables = []
    for n in (1, 7):
        snap = Snapshot(devices=[switch(n)])
        doc = build_summary(snap)
        assert _device_pages(snap, doc) == [1], n
        vlan_tables.append([b for b in doc.blocks if getattr(b, "columns", None) == ["VLAN", "Name", "Subnet / gateway"]])
    (table,) = vlan_tables[0]
    assert 10 <= len(table.rows) < 59 and table.note == f"+{59 - len(table.rows)} more in the As-Built"
    assert ("Uptime", "214d 6h") in next(b for b in doc.blocks if getattr(b, "title", None) == "At a glance").items
    if not vlan_tables[1]:
        assert any(isinstance(b, Paragraph) and b.text == "59 VLANs: see the As-Built." for b in doc.blocks)


def test_word_page_breaks_are_folded_into_the_next_heading(snap, tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    from sextant.render.docx import render_docx

    body = Document(render_docx(build_summary(snap), tmp_path / "s.docx")).element.body
    breaks = body.findall(f".//{qn('w:br')}[@{qn('w:type')}='page']")
    before = body.findall(f".//{qn('w:pageBreakBefore')}")
    assert len(before) >= len(snap.devices)
    for br in breaks:  # what is left is followed by a table, where a paragraph property cannot carry it
        assert br.getparent().getparent().getnext().tag == qn("w:tbl")


def test_table_rows_match_the_columns():
    from sextant.doc import Cell, Table

    t = Table(["a", "b", "c"], [["1"], ["1", "2", "3", Cell("4", "error"), "5"]])
    assert t.rows == [["1", "", ""], ["1", "2", "3 · 4 · 5"]]


def test_neighbor_table_lists_each_cable_once():
    """Both ends report the cable with different names: FQDN vs short name, long vs short port names, and
    one end that does not know the far port."""
    from sextant.model import Device, Neighbor, Snapshot

    a = Device(hostname="sw-a", role="switch", neighbors=[
        Neighbor(local_port="TenGigabitEthernet1/1/1", remote_device="sw-b.corp.local", remote_port="Gi1/0/48"),
        Neighbor(local_port="Gi1/0/2", remote_device="sw-c", remote_port="Gi0/1", protocol="cdp")])
    b = Device(hostname="SW-B", role="switch", neighbors=[
        Neighbor(local_port="GigabitEthernet1/0/48", remote_device="sw-a.corp.local", remote_port="Te1/1/1")])
    c = Device(hostname="sw-c", role="switch", neighbors=[Neighbor(local_port="Gi0/1", remote_device="sw-a")])
    doc = build_asbuilt(Snapshot(devices=[a, b, c]))
    (table,) = [blk for blk in doc.blocks if isinstance(getattr(blk, "columns", None), list)
                and "Neighbor IP" in blk.columns]
    assert sorted((r[0], r[1]) for r in table.rows) == [("sw-a", "Gi1/0/2"), ("sw-a", "Te1/1/1")]


def test_lag_vlans_column_matches_the_interface_table():
    from sextant.model import Device, Interface, PortMode, Snapshot

    d = Device(hostname="sw", role="switch", interfaces=[
        Interface(name="Po1", kind="lag", mode=PortMode.trunk, lag_members=["Gi1/0/1"]),
        Interface(name="Po2", kind="lag", mode=PortMode.trunk, trunk_vlans="10,20", native_vlan=99),
        Interface(name="Po3", kind="lag", mode=PortMode.access, access_vlan=30),
        Interface(name="Trk1", kind="lag", mode=PortMode.trunk, trunk_vlans="10", native_vlan=0)])
    doc = build_asbuilt(Snapshot(devices=[d]))
    (lag,) = [blk for blk in doc.blocks if getattr(blk, "title", None) == "Link aggregation"]
    assert [r[-1] for r in lag.rows] == ["all (native 1)", "10,20 (native 99)", "30", "10 (native none, all tagged)"]
