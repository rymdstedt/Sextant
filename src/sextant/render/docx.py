"""Render a doc.Document to Word (.docx) in the active theme (python-docx)."""

from __future__ import annotations

from contextvars import ContextVar
from dataclasses import dataclass
from pathlib import Path

from docx import Document as new_docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from sextant import doc as d
from sextant.theme import Theme, load_theme

STATUS_MARK = {"success": "✓", "warning": "◆", "error": "▲", "info": "●", "neutral": "●"}


@dataclass(frozen=True)
class Palette:
    """Theme colours as Word hex strings ("336699")."""

    font: str
    mono: str
    primary: str
    highlight: str
    text: str
    on_primary: str
    on_primary_muted: str
    gray_50: str
    gray_100: str
    gray_200: str
    gray_600: str
    gray_700: str
    status: dict[str, tuple[str, str]]  # status -> (text colour, background)
    border: dict[str, str]  # status -> strong border colour (KPI tiles, callouts)

    @classmethod
    def from_theme(cls, theme: Theme) -> Palette:
        c = theme.colors
        h = c.hex
        return cls(
            font=theme.fonts.family, mono=theme.fonts.mono, primary=h("primary"), highlight=h("highlight"),
            text=h("text"), on_primary=h("on_primary"), on_primary_muted=h("on_primary_muted"),
            gray_50=h("gray_50"), gray_100=h("gray_100"), gray_200=h("gray_200"), gray_600=h("gray_600"),
            gray_700=h("gray_700"),
            status={"success": (h("success"), h("success_bg")), "warning": (h("warning"), h("warning_bg")),
                    "error": (h("error"), h("error_bg")), "info": (h("info"), h("info_bg")),
                    "neutral": (h("gray_700"), h("gray_100"))},
            border={"success": h("success"), "warning": h("warning_border"), "error": h("error"),
                    "info": h("info"), "neutral": h("gray_500")},
        )


_palette: ContextVar[Palette] = ContextVar("docx_palette")


def P() -> Palette:  # noqa: N802 - short accessor used throughout
    try:
        return _palette.get()
    except LookupError:
        return Palette.from_theme(load_theme())


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# --------------------------------------------------------------------------- low-level XML helpers


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_borders(cell, **edges: tuple[str, int]) -> None:
    """edges: top=("336699", 18) -> colour, size in eighths of a point. Use ("nil", 0) to hide."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, (color, size) in edges.items():
        el = OxmlElement(f"w:{edge}")
        if color == "nil":
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        borders.append(el)


def _table_borders(table, inside_h: str | None = "default", outer: bool = False) -> None:
    if inside_h == "default":
        inside_h = P().gray_200
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        show = (edge == "insideH" and inside_h) or (edge == "bottom" and inside_h) or (outer and edge in
                                                                                       ("top", "left", "right"))
        if show:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), inside_h or P().gray_200)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell_margins(table, top=40, bottom=40, left=80, right=80) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)
    if placeholder:
        t = OxmlElement("w:t")
        t.text = placeholder
        run._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _emu_to_dxa(emu: int) -> int:
    return int(emu / 635)


def _set_widths(table, widths: list[int]) -> None:
    """Fixed layout with explicit grid/cell widths (EMU). Word and LibreOffice both honour this."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(_emu_to_dxa(sum(widths))))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths, strict=False):
        col.set(qn("w:w"), str(_emu_to_dxa(w)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths, strict=False):
            cell.width = w


def _set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None,
              weight_font: str | None = None) -> None:
    run.font.name = weight_font or P().font
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), weight_font or P().font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------- renderer


class DocxRenderer:
    def __init__(self, doc: d.Document, theme: Theme):
        self.doc = doc
        self.theme = theme
        self.w = new_docx()
        self._styles()
        sec = self.w.sections[0]
        if doc.landscape:
            sec.orientation = WD_ORIENT.landscape
            sec.page_width, sec.page_height = Mm(297), Mm(210)
        else:
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(16)
        sec.top_margin, sec.bottom_margin = Mm(22), Mm(18)
        sec.header_distance = sec.footer_distance = Mm(8)
        sec.different_first_page_header_footer = True
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        self._header_footer(sec)
        self._first_h1 = True
        settings = self.w.settings.element
        upd = OxmlElement("w:updateFields")  # refresh the TOC when opened in Word
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    # ---- styles
    def _styles(self) -> None:
        st = self.w.styles
        normal = st["Normal"]
        normal.font.name = P().font
        normal.font.size = Pt(self.theme.fonts.word_body_pt)
        normal.font.color.rgb = _rgb(P().text)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), P().font)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        sizes = (list(self.theme.fonts.word_heading_pt) + [20, 18, 14, 12][len(self.theme.fonts.word_heading_pt):])[:4]
        for lvl, size in enumerate(sizes, start=1):  # heading scale from the theme
            h = st[f"Heading {lvl}"]
            h.font.name = P().font
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = _rgb(P().primary)
            rfonts = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(a), P().font)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rfonts.get(qn(a)) is not None:
                    del rfonts.attrib[qn(a)]
            h.paragraph_format.space_before = Pt({1: 0, 2: 14, 3: 10, 4: 8}[lvl])
            h.paragraph_format.space_after = Pt({1: 8, 2: 5, 3: 4, 4: 3}[lvl])
            h.paragraph_format.keep_with_next = True

    def _header_footer(self, sec) -> None:
        hp = sec.header.paragraphs[0]
        tbl = sec.header.add_table(1, 2, self.text_width)
        _table_borders(tbl, inside_h=None)
        _set_widths(tbl, [int(self.text_width * 0.7), int(self.text_width * 0.3)])
        left, right = tbl.rows[0].cells
        left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        r = left.paragraphs[0].add_run(self.doc.header_text)
        _set_font(r, 7.5, color=P().gray_600)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        logo = self.theme.asset(self.theme.logos.color)
        if logo:
            rp.add_run().add_picture(str(logo), width=Mm(22))
        elif self.theme.company:
            _set_font(rp.add_run(self.theme.company), 8, True, P().primary)
        hp._p.getparent().remove(hp._p)

        fp = sec.footer.paragraphs[0]
        ft = sec.footer.add_table(1, 2, self.text_width)
        _table_borders(ft, inside_h=None)
        _set_widths(ft, [int(self.text_width * 0.6), int(self.text_width * 0.4)])
        left, right = ft.rows[0].cells
        _set_font(left.paragraphs[0].add_run(self.doc.footer_text), 7.5, color=P().gray_600)
        rp = right.paragraphs[0]
        footer_style = self.w.styles["Footer"]
        footer_style.font.size = Pt(7.5)
        footer_style.font.color.rgb = _rgb(P().gray_600)
        rp.style = footer_style
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_font(rp.add_run("Page "), 7.5, color=P().gray_600)
        _field(rp, "PAGE", "1")
        _set_font(rp.add_run(" of "), 7.5, color=P().gray_600)
        _field(rp, "NUMPAGES", "1")
        for r in rp.runs:
            _set_font(r, 7.5, color=P().gray_600)
        fp._p.getparent().remove(fp._p)

    # ---- blocks
    def render(self, out: Path) -> Path:
        for b in self.doc.blocks:
            getattr(self, f"_{type(b).__name__.lower()}")(b)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.w.core_properties.title = self.doc.title
        self.w.core_properties.author = self.theme.company or self.theme.product
        self.w.save(out)
        return out

    def _para(self, text: str = "", size: float | None = None, bold: bool | None = None, color: str | None = None,
              after: float | None = None):
        p = self.w.add_paragraph()
        if text:
            _set_font(p.add_run(text), size, bold, color)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def _cover(self, b: d.Cover) -> None:
        p = self.w.paragraphs[0] if self.w.paragraphs else self._para()
        p.paragraph_format.space_before = Mm(4)
        logo = self.theme.asset(self.theme.logos.color)
        if logo:
            p.add_run().add_picture(str(logo), width=Mm(62))
        elif self.theme.company:
            _set_font(p.add_run(self.theme.company), 20, True, P().primary)
        self._para(after=0).paragraph_format.space_before = Mm(55 if not self.doc.landscape else 20)
        t = self.w.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, inside_h=None)
        _set_widths(t, [int(self.text_width)])
        c = t.rows[0].cells[0]
        _shade(c, P().primary)
        _cell_borders(c, top=(P().highlight, 36))
        _cell_margins(t, 500, 600, 500, 500)
        cp = c.paragraphs[0]
        _set_font(cp.add_run((b.customer or "").upper()), 11, True, P().on_primary)
        cp.paragraph_format.space_after = Pt(6)
        tp = c.add_paragraph()
        _set_font(tp.add_run(b.title), 32, True, P().on_primary)
        tp.paragraph_format.space_after = Pt(4)
        if b.subtitle:
            sp = c.add_paragraph()
            _set_font(sp.add_run(b.subtitle), 16, True, P().on_primary)
            sp.paragraph_format.space_after = Pt(22)
        for k, v in (("Site", b.site), ("Date", b.date), ("Version", b.version),
                     ("Classification", b.classification)):
            if v:
                mp = c.add_paragraph()
                mp.paragraph_format.space_after = Pt(2)
                mp.paragraph_format.tab_stops.add_tab_stop(Mm(36))
                _set_font(mp.add_run(f"{k}\t"), 10.5, True, P().on_primary_muted)
                _set_font(mp.add_run(v), 10.5, False, P().on_primary)
        self._pagebreak(None)

    def _heading(self, b: d.Heading) -> None:
        if b.level == 1 and not self._first_h1:
            self._pagebreak(None)
        self._first_h1 = False if b.level == 1 else self._first_h1
        h = self.w.add_heading(b.text, level=min(b.level, 4))
        if b.level == 1:
            p_pr = h._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "4"), ("w:color", P().primary)):
                bottom.set(qn(k), v)
            bdr.append(bottom)
            p_pr.append(bdr)
        if not b.toc:  # keep out of the Word TOC: use outline level body text
            p_pr = h._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), "9")
            p_pr.append(ol)

    def _paragraph(self, b: d.Paragraph) -> None:
        style = {"body": (None, None, None), "ingress": (11, True, P().gray_700), "muted": (None, None, P().gray_600),
                 "footnote": (8, None, P().gray_600)}[b.style]
        self._para(b.text, *style)

    def _facts(self, b: d.Facts) -> None:
        if b.title:
            _keep_with_next(self._para(b.title, 10.5, True, P().primary, after=3))
        cols = b.columns
        rows = [b.items[i:i + cols] for i in range(0, len(b.items), cols)]
        t = self.w.add_table(rows=len(rows), cols=2 * cols)
        _table_borders(t, inside_h=P().gray_100)
        _cell_margins(t, 30, 30, 60, 60)
        kw, vw = (0.19, 0.31) if cols == 2 else (0.28, 0.72)
        _set_widths(t, [int(self.text_width * x) for x in [kw, vw] * cols])
        for r_i, row in enumerate(rows):
            _no_split(t.rows[r_i])
            for c_i, (k, v) in enumerate(row):
                kc, vc = t.rows[r_i].cells[2 * c_i], t.rows[r_i].cells[2 * c_i + 1]
                _set_font(kc.paragraphs[0].add_run(k), 8.5, True, P().gray_600)
                _set_font(vc.paragraphs[0].add_run(v), 8.5)
                for c in (kc, vc):
                    c.paragraphs[0].paragraph_format.space_after = Pt(0)
        self._para(after=2)

    def _cell_text(self, cell, value, size: float, align_right: bool = False) -> None:
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if align_right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if isinstance(value, d.Cell):
            status = value.status or "neutral"
            fg, bg = P().status[status]
            _set_font(p.add_run(f"{STATUS_MARK[status]} {value.text}"), size, True, fg)
            _shade(cell, bg)
        else:
            _set_font(p.add_run(str(value)), size)

    def _table(self, b: d.Table) -> None:
        if b.title:
            _keep_with_next(self._para(b.title, 10, True, P().primary, after=3))
        size = 7.5 if b.compact else 8.5
        t = self.w.add_table(rows=1 + len(b.rows), cols=len(b.columns))
        t.autofit = False
        _table_borders(t)
        _cell_margins(t, 30, 30, 60, 60) if b.compact else _cell_margins(t)
        widths = b.widths or [1.0] * len(b.columns)
        total = sum(widths)
        col_w = [int(self.text_width * w / total) for w in widths]
        _set_widths(t, col_w)
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, col in enumerate(b.columns):
            c = hdr.cells[i]
            c.width = col_w[i]
            _shade(c, P().gray_50)
            _cell_borders(c, bottom=(P().gray_200, 8))
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i in b.numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_font(p.add_run(col), size, True, P().primary)
        for r_i, row in enumerate(b.rows, start=1):
            tr = t.rows[r_i]
            _no_split(tr)
            for c_i, value in enumerate(row[:len(b.columns)]):
                c = tr.cells[c_i]
                c.width = col_w[c_i]
                self._cell_text(c, value, size, c_i in b.numeric)
        if b.note:
            self._para(b.note, 7.5, None, P().gray_600)
        else:
            self._para(after=2)

    def _kpirow(self, b: d.KpiRow) -> None:
        n = len(b.items)
        if not n:
            return
        t = self.w.add_table(rows=1, cols=2 * n - 1)  # spacer columns between tiles
        t.autofit = False
        _table_borders(t, inside_h=None)
        _cell_margins(t, 80, 80, 110, 80)
        gap = int(Mm(3))
        tile = int((self.text_width - gap * (n - 1)) / n)
        _set_widths(t, [tile if i % 2 == 0 else gap for i in range(2 * n - 1)])
        _no_split(t.rows[0])
        for i, k in enumerate(b.items):
            c = t.rows[0].cells[2 * i]
            c.width = tile
            if i:
                t.rows[0].cells[2 * i - 1].width = gap
            edge = (P().gray_200, 4)
            _cell_borders(c, top=(P().border.get(k.status or "", P().primary), 24), left=edge, right=edge, bottom=edge)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            _set_font(p.add_run(k.label.upper()), 7, True, P().gray_600)
            vp = c.add_paragraph()
            vp.paragraph_format.space_after = Pt(0)
            _set_font(vp.add_run(k.value), 18 if len(k.value) < 10 else 12, True, P().primary)
            if k.sub:
                sp = c.add_paragraph()
                sp.paragraph_format.space_after = Pt(0)
                _set_font(sp.add_run(k.sub), 7, None, P().gray_600)
            if k.status in ("error", "warning", "success"):
                label = {"error": "▲ Action needed", "warning": "◆ Review", "success": "✓ OK"}[k.status]
                stp = c.add_paragraph()
                stp.paragraph_format.space_after = Pt(0)
                _set_font(stp.add_run(label), 7, True, P().status[k.status][0])
        self._para(after=4)

    def _findings(self, b: d.Findings) -> None:
        if b.title:
            _keep_with_next(self._para(b.title, 11, True, P().primary, after=3))
        if not b.items:
            self._para(f"✓ {b.empty_text}", None, True, P().status["success"][0])
            return
        t = self.w.add_table(rows=len(b.items), cols=2)
        _set_widths(t, [int(Mm(22)), int(self.text_width - Mm(22))])
        _table_borders(t)
        _cell_margins(t, 50, 50, 60, 60)
        for i, f in enumerate(b.items):
            row = t.rows[i]
            _no_split(row)
            sev_c, body = row.cells
            sev_c.width, body.width = int(Mm(22)), int(self.text_width - Mm(22))
            status = d.SEVERITY_STATUS[f.severity]
            fg, bg = P().status[status]
            p = sev_c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _set_font(p.add_run(f"{STATUS_MARK[status]} {d.SEVERITY_LABEL[f.severity]}"), 8, True, fg)
            _shade(sev_c, bg)
            bp = body.paragraphs[0]
            bp.paragraph_format.space_after = Pt(1)
            if f.device:
                _set_font(bp.add_run(f"{f.device}: "), 9, True, P().gray_600)
            _set_font(bp.add_run(f.title), 9, True)
            if f.detail and not b.compact:
                dp = body.add_paragraph()
                dp.paragraph_format.space_after = Pt(1)
                _set_font(dp.add_run(f.detail), 8.5, None, P().gray_700)
            if f.fix:
                fp = body.add_paragraph()
                fp.paragraph_format.space_after = Pt(0)
                _set_font(fp.add_run("Fix: "), 8.5, True, P().primary)
                _set_font(fp.add_run(f.fix), 8.5, None, P().gray_700)
        self._para(after=2)

    def _code(self, b: d.Code) -> None:
        if b.title:
            self._para(b.title, 10, True, P().primary)
        t = self.w.add_table(rows=1, cols=1)
        _table_borders(t, inside_h=None)
        _set_widths(t, [int(self.text_width)])
        c = t.rows[0].cells[0]
        _shade(c, P().gray_50)
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        lines = b.text.splitlines()
        for i, line in enumerate(lines):
            r = p.add_run(line)
            _set_font(r, 6.5, weight_font=P().mono)
            if i < len(lines) - 1:
                r.add_break()
        self._para(after=2)

    def _callout(self, b: d.Callout) -> None:
        t = self.w.add_table(rows=1, cols=1)
        _table_borders(t, inside_h=None)
        _cell_margins(t, 80, 80, 140, 100)
        _set_widths(t, [int(self.text_width)])
        c = t.rows[0].cells[0]
        fg, bg = P().status[b.status]
        _shade(c, bg)
        _cell_borders(c, left=(P().border.get(b.status, fg), 24))
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if b.title:
            _set_font(p.add_run(f"{b.title}: "), 9, True, P().primary)
        _set_font(p.add_run(b.text), 9)
        self._para(after=2)

    def _deviceheader(self, b: d.DeviceHeader) -> None:
        t = self.w.add_table(rows=1, cols=2)
        t.autofit = False
        _table_borders(t, inside_h=None)
        _cell_margins(t, 90, 90, 140, 120)
        _set_widths(t, [int(self.text_width * 0.75), int(self.text_width * 0.25)])
        left, right = t.rows[0].cells
        for c in (left, right):
            _shade(c, P().gray_50)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _cell_borders(left, left=(P().primary, 36))
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _set_font(p.add_run(b.role.upper()), 7.5, True, P().gray_600)
        lp = left.add_paragraph()
        lp.paragraph_format.space_after = Pt(0)
        _set_font(lp.add_run(b.line), 9.5, True, P().gray_700)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        fg, _bg = P().status[b.status]
        _set_font(rp.add_run(f"{STATUS_MARK[b.status]} {b.status_label}"), 11, True, fg)
        self._para(after=4)

    def _diagram(self, b: d.Diagram) -> None:
        from io import BytesIO

        from sextant.render.diagram import svg_to_png
        from sextant.render.stp_diagram import block_svg

        svg, _ = block_svg(b, self.theme)
        png = svg_to_png(svg)
        if png is None:
            self._callout(d.Callout("The topology diagram needs cairosvg; see the PDF or HTML version.", "info"))
            return
        p = self.w.add_paragraph()
        p.paragraph_format.keep_with_next = bool(b.caption)
        p.add_run().add_picture(BytesIO(png), width=self.text_width)
        if b.caption:
            self._para(b.caption, 7.5, None, P().gray_600)

    def _pagebreak(self, _b) -> None:
        p = self.w.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)
        self._first_h1 = True

    def _toc(self, b: d.Toc) -> None:
        h = self.w.add_heading(b.title, level=1)
        p_pr = h._p.get_or_add_pPr()
        ol = OxmlElement("w:outlineLvl")
        ol.set(qn("w:val"), "9")
        p_pr.append(ol)
        entries = [x for x in self.doc.headings() if x.level <= 2] or [d.Heading("Contents")]
        paras = []
        for i, entry in enumerate(entries):
            p = self.w.add_paragraph()
            p.paragraph_format.space_after = Pt(2 if entry.level == 2 else 1)
            p.paragraph_format.space_before = Pt(6 if entry.level == 1 else 0)
            if entry.level == 2:
                p.paragraph_format.left_indent = Mm(6)
            if i == 0:  # field begin + instruction + separator before the first entry
                run = p.add_run()
                for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
                    el = OxmlElement(tag)
                    if attr:
                        el.set(qn("w:fldCharType"), attr)
                    else:
                        el.set(qn("xml:space"), "preserve")
                        el.text = 'TOC \\o "1-2" \\h \\z \\u'
                    run._r.append(el)
            # pre-filled result: visible before Word refreshes the field (which then adds page numbers)
            _set_font(p.add_run(entry.text), 9 if entry.level == 2 else 10, entry.level == 1,
                      P().primary if entry.level == 1 else None)
            paras.append(p)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paras[-1].add_run()._r.append(end)
        self._pagebreak(None)

def render_docx(doc: d.Document, out: Path, theme: Theme | None = None) -> Path:
    theme = theme or load_theme()
    token = _palette.set(Palette.from_theme(theme))
    try:
        return DocxRenderer(doc, theme).render(out)
    finally:
        _palette.reset(token)
